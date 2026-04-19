#!/usr/bin/env python3
"""
Genera documento Word de Arquitectura Cerebral SOUL — en Español
================================================================
Crea un .docx completo explicando toda la arquitectura cognitiva SOUL
con diagramas generados por matplotlib embebidos como imagenes.

Autor: JARVIS (Equipo SEAL)
Fecha: 2026-04-09
"""

import os
import io
import math
import textwrap
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import numpy as np

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

# ── Rutas de salida ──
OUT_DIR = Path(__file__).parent
IMG_DIR = OUT_DIR / "_img_cache"
IMG_DIR.mkdir(exist_ok=True)
DOCX_PATH = OUT_DIR / "SOUL_Arquitectura_Cerebral_v1.docx"

# ── Paleta de colores ──
COLORS = {
    "primary": "#1a237e",
    "secondary": "#0d47a1",
    "accent": "#00bcd4",
    "warm": "#ff6f00",
    "success": "#2e7d32",
    "danger": "#c62828",
    "purple": "#6a1b9a",
    "bg": "#f5f5f5",
    "text": "#212121",
    "light_text": "#ffffff",
    "identity": "#7b1fa2",
    "memory": "#1565c0",
    "emotion": "#c62828",
    "instinct": "#e65100",
    "reflection": "#00695c",
    "graph": "#283593",
    "rules": "#4e342e",
}


def save_fig(fig, name: str) -> str:
    path = IMG_DIR / f"{name}.png"
    fig.savefig(path, dpi=200, bbox_inches='tight', facecolor='white', edgecolor='none')
    plt.close(fig)
    return str(path)


# ════════════════════════════════════════════════════════════════
# GENERADORES DE DIAGRAMAS
# ════════════════════════════════════════════════════════════════

def diagram_01_overview():
    """Vista general de la arquitectura SOUL."""
    fig, ax = plt.subplots(figsize=(12, 8))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 8)
    ax.axis('off')
    ax.set_facecolor('white')

    ax.text(6, 7.6, "SOUL — Vista General de la Arquitectura Cognitiva", fontsize=15,
            fontweight='bold', ha='center', color=COLORS["primary"])

    # Caja central del cerebro
    brain = FancyBboxPatch((3.5, 3.2), 5, 3.2, boxstyle="round,pad=0.15",
                            facecolor="#e8eaf6", edgecolor=COLORS["primary"], linewidth=2.5)
    ax.add_patch(brain)
    ax.text(6, 5.9, "Motor SOUL", fontsize=14, fontweight='bold',
            ha='center', color=COLORS["primary"])

    modules = [
        (4.2, 4.8, "Identidad\n(OCEAN)", COLORS["identity"]),
        (6.0, 4.8, "Almacen\nMemoria", COLORS["memory"]),
        (7.8, 4.8, "Motor\nReflexion", COLORS["reflection"]),
        (4.2, 3.6, "Instintos", COLORS["instinct"]),
        (6.0, 3.6, "Motor\nReglas", COLORS["rules"]),
        (7.8, 3.6, "Conectoma\n(Grafo)", COLORS["graph"]),
    ]

    for x, y, label, color in modules:
        box = FancyBboxPatch((x-0.7, y-0.4), 1.4, 0.8, boxstyle="round,pad=0.08",
                              facecolor=color, edgecolor='white', linewidth=1, alpha=0.85)
        ax.add_patch(box)
        ax.text(x, y, label, fontsize=7.5, fontweight='bold',
                ha='center', va='center', color='white')

    externals = [
        (1.2, 5.5, "PostgreSQL\n(Memorias)", "#1565c0"),
        (1.2, 3.8, "Neo4j\n(Grafo)", "#283593"),
        (10.8, 5.5, "Qdrant\n(Vectores)", "#00695c"),
        (10.8, 3.8, "Ollama\n(LLM)", "#e65100"),
        (6, 1.5, "Agente LLM (Claude, GPT, Ollama, ...)", "#37474f"),
    ]

    for x, y, label, color in externals:
        w = 1.8 if y > 2 else 4
        box = FancyBboxPatch((x-w/2, y-0.4), w, 0.8, boxstyle="round,pad=0.08",
                              facecolor=color, edgecolor='white', linewidth=1)
        ax.add_patch(box)
        ax.text(x, y, label, fontsize=8, fontweight='bold',
                ha='center', va='center', color='white')

    arrows = [
        (2.2, 5.5, 3.5, 5.5), (2.2, 3.8, 3.5, 3.8),
        (9.7, 5.5, 8.5, 5.5), (9.7, 3.8, 8.5, 3.8),
        (6, 3.2, 6, 2.0),
    ]
    for x1, y1, x2, y2 in arrows:
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="<->", color="#546e7a", lw=1.5))

    ax.text(1.2, 2.5, "Almacenamiento\nBackend", fontsize=7, ha='center', color='#546e7a', style='italic')
    ax.text(10.8, 2.5, "Servicios IA", fontsize=7, ha='center', color='#546e7a', style='italic')
    ax.text(6, 2.65, "boot_context()\nsnapshot()", fontsize=7, ha='center', color='#546e7a', style='italic')

    return save_fig(fig, "01_overview")


def diagram_02_memory_flow():
    """Ciclo de vida de la memoria."""
    fig, ax = plt.subplots(figsize=(12, 5))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 5)
    ax.axis('off')

    ax.text(6, 4.7, "Ciclo de Vida de la Memoria", fontsize=14, fontweight='bold',
            ha='center', color=COLORS["primary"])

    stages = [
        (1.5, 2.5, "GUARDAR", "El agente guarda\nmemoria con\nimportancia 1-10", "#1565c0"),
        (4.0, 2.5, "VECTORIZAR", "TF-IDF o\nSentenceTransformer\ngenera vector", "#00695c"),
        (6.5, 2.5, "INDEXAR", "PostgreSQL +\nQdrant escritura\ndual", "#283593"),
        (9.0, 2.5, "BUSCAR", "Busqueda hibrida:\ntexto + vector\n+ scoring", "#6a1b9a"),
        (11.0, 2.5, "DECAER", "decaimiento\ntemporal reduce\npuntaje", "#c62828"),
    ]

    for x, y, title, desc, color in stages:
        box = FancyBboxPatch((x-0.8, y-0.7), 1.6, 1.4, boxstyle="round,pad=0.1",
                              facecolor=color, edgecolor='white', linewidth=1.5)
        ax.add_patch(box)
        ax.text(x, y+0.25, title, fontsize=9, fontweight='bold',
                ha='center', va='center', color='white')
        ax.text(x, y-0.25, desc, fontsize=6.5, ha='center', va='center', color='#e0e0e0')

    for i in range(len(stages)-1):
        x1 = stages[i][0] + 0.85
        x2 = stages[i+1][0] - 0.85
        ax.annotate("", xy=(x2, 2.5), xytext=(x1, 2.5),
                    arrowprops=dict(arrowstyle="->", color="#37474f", lw=2))

    box = FancyBboxPatch((2.5, 0.3), 7, 1.0, boxstyle="round,pad=0.1",
                          facecolor="#fff3e0", edgecolor="#ff6f00", linewidth=1.5)
    ax.add_patch(box)
    ax.text(6, 0.95, "Formula de Puntuacion", fontsize=9, fontweight='bold',
            ha='center', color="#e65100")
    ax.text(6, 0.55, "puntaje = importancia * frecuencia_acceso * decaimiento_temporal(horas, vida_media=168h)",
            fontsize=7.5, ha='center', color="#424242", family='monospace')

    return save_fig(fig, "02_memory_flow")


def diagram_03_ocean():
    """Visualizacion del modelo de personalidad OCEAN."""
    fig, axes = plt.subplots(1, 2, figsize=(12, 5))

    ocean_jarvis = {"Apertura": 0.782, "Responsabilidad": 0.948,
                    "Extraversion": 0.392, "Amabilidad": 0.661, "Neuroticismo": 0.115}
    colors_bars = [COLORS["identity"], COLORS["success"], COLORS["accent"],
                   "#ff6f00", COLORS["danger"]]

    ax1 = axes[0]
    bars = ax1.barh(list(ocean_jarvis.keys()), list(ocean_jarvis.values()),
                    color=colors_bars, edgecolor='white', height=0.6)
    ax1.set_xlim(0, 1.0)
    ax1.set_title("JARVIS — Perfil OCEAN", fontsize=12, fontweight='bold', color=COLORS["primary"])
    for bar, val in zip(bars, ocean_jarvis.values()):
        ax1.text(val + 0.02, bar.get_y() + bar.get_height()/2,
                f"{val:.3f}", va='center', fontsize=9, fontweight='bold')
    ax1.spines['top'].set_visible(False)
    ax1.spines['right'].set_visible(False)

    ax2 = axes[1]
    ax2.remove()
    ax2 = fig.add_subplot(122, polar=True)
    labels = list(ocean_jarvis.keys())
    values = list(ocean_jarvis.values())
    values += values[:1]
    angles = np.linspace(0, 2*np.pi, len(labels), endpoint=False).tolist()
    angles += angles[:1]

    ax2.fill(angles, values, color=COLORS["identity"], alpha=0.25)
    ax2.plot(angles, values, color=COLORS["identity"], linewidth=2)
    ax2.set_xticks(angles[:-1])
    ax2.set_xticklabels([l[0] for l in labels], fontsize=11, fontweight='bold')
    ax2.set_ylim(0, 1.0)
    ax2.set_title("Radar OCEAN", fontsize=12, fontweight='bold', color=COLORS["primary"], pad=20)

    fig.tight_layout()
    return save_fig(fig, "03_ocean")


def diagram_04_scoring_decay():
    """Curva de decaimiento temporal."""
    fig, ax = plt.subplots(figsize=(10, 5))

    hours = np.linspace(0, 720, 500)
    half_life = 168

    decay = np.exp(-0.693 * hours / half_life)
    emotional_factor = 1 / (1 + 0.8*0.5 + 0.7*0.3)
    decay_emotional = np.exp(-0.693 * hours / half_life * emotional_factor)
    decay_important = np.exp(-0.693 * hours / (half_life * 1.5))

    ax.plot(hours/24, decay, color=COLORS["memory"], linewidth=2.5, label="Decaimiento estandar (imp=5)")
    ax.plot(hours/24, decay_emotional, color=COLORS["emotion"], linewidth=2.5,
            label="Memoria emocional (valencia=0.8)", linestyle='--')
    ax.plot(hours/24, decay_important, color=COLORS["success"], linewidth=2.5,
            label="Memoria importante (imp=9)", linestyle='-.')

    ax.axhline(y=0.5, color='gray', linestyle=':', alpha=0.5)
    ax.axvline(x=7, color='gray', linestyle=':', alpha=0.5)
    ax.text(7.5, 0.52, "7 dias\n(vida media)", fontsize=8, color='gray')

    ax.set_xlabel("Dias", fontsize=11)
    ax.set_ylabel("Puntaje de Retencion", fontsize=11)
    ax.set_title("Decaimiento Temporal — Retencion de Memoria en el Tiempo", fontsize=13,
                 fontweight='bold', color=COLORS["primary"])
    ax.legend(fontsize=9)
    ax.set_xlim(0, 30)
    ax.set_ylim(0, 1.05)
    ax.grid(True, alpha=0.2)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)

    return save_fig(fig, "04_scoring_decay")


def diagram_05_connectome():
    """Visualizacion del grafo de conocimiento / conectoma."""
    fig, ax = plt.subplots(figsize=(10, 8))
    ax.set_xlim(-1.5, 1.5)
    ax.set_ylim(-1.5, 1.5)
    ax.axis('off')
    ax.set_facecolor('white')

    ax.text(0, 1.4, "Conectoma SOUL — Grafo de Conocimiento", fontsize=14,
            fontweight='bold', ha='center', color=COLORS["primary"])

    n_nodes = 12
    node_data = [
        ("Memoria\n#142", COLORS["memory"]),
        ("Memoria\n#89", COLORS["memory"]),
        ("Memoria\n#201", COLORS["memory"]),
        ("Memoria\n#55", COLORS["emotion"]),
        ("Memoria\n#167", COLORS["memory"]),
        ("Entidad:\nWilliam", COLORS["identity"]),
        ("Entidad:\nADA", COLORS["identity"]),
        ("Entidad:\nPostgreSQL", COLORS["instinct"]),
        ("Memoria\n#312", COLORS["memory"]),
        ("Correccion\n#78", COLORS["danger"]),
        ("Memoria\n#445", COLORS["reflection"]),
        ("Entidad:\nSOUL", COLORS["instinct"]),
    ]

    positions = []
    for i in range(n_nodes):
        angle = 2 * np.pi * i / n_nodes - np.pi/2
        r = 0.95
        x = r * np.cos(angle)
        y = r * np.sin(angle)
        positions.append((x, y))

    edges = [
        (0, 1, "EXCITA", COLORS["success"], 0.85),
        (0, 2, "EXCITA", COLORS["success"], 0.72),
        (1, 3, "EXCITA", COLORS["success"], 0.78),
        (2, 9, "INHIBE", COLORS["danger"], 0.91),
        (3, 4, "EXCITA", COLORS["success"], 0.71),
        (0, 5, "MENCIONA", COLORS["identity"], 0.0),
        (1, 6, "MENCIONA", COLORS["identity"], 0.0),
        (4, 7, "MENCIONA", COLORS["instinct"], 0.0),
        (8, 11, "MENCIONA", COLORS["instinct"], 0.0),
        (5, 6, "EXCITA", COLORS["success"], 0.80),
        (9, 10, "INHIBE", COLORS["danger"], 0.88),
        (8, 0, "EXCITA", COLORS["success"], 0.74),
        (10, 11, "MENCIONA", COLORS["instinct"], 0.0),
    ]

    for src, dst, rel, color, weight in edges:
        x1, y1 = positions[src]
        x2, y2 = positions[dst]
        style = '--' if rel == "INHIBE" else ('-' if rel == "EXCITA" else ':')
        lw = 1.5 if weight > 0.8 else 1.0
        ax.plot([x1, x2], [y1, y2], color=color, linestyle=style, linewidth=lw, alpha=0.5)

    for i, ((x, y), (label, color)) in enumerate(zip(positions, node_data)):
        circle = plt.Circle((x, y), 0.15, facecolor=color, edgecolor='white', linewidth=1.5, zorder=5)
        ax.add_patch(circle)
        ax.text(x, y, label, fontsize=5.5, fontweight='bold',
                ha='center', va='center', color='white', zorder=6)

    legend_items = [
        mpatches.Patch(color=COLORS["success"], label="EXCITA (refuerza)"),
        mpatches.Patch(color=COLORS["danger"], label="INHIBE (contradice)"),
        mpatches.Patch(color=COLORS["identity"], label="MENCIONA (enlace entidad)"),
    ]
    ax.legend(handles=legend_items, loc='lower right', fontsize=8)

    return save_fig(fig, "05_connectome")


def diagram_06_reflection_loop():
    """Ciclo de auto-reflexion y monologo interno."""
    fig, ax = plt.subplots(figsize=(10, 6))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis('off')

    ax.text(5, 5.7, "Ciclo de Auto-Reflexion", fontsize=14, fontweight='bold',
            ha='center', color=COLORS["primary"])

    steps = [
        (2, 4, "Experiencia\n(interaccion)", COLORS["memory"]),
        (5, 4, "self_reflect()\n(pensamiento\n+ emocion)", COLORS["reflection"]),
        (8, 4, "inner_thoughts\n(monologo\nprivado)", COLORS["identity"]),
        (8, 2, "Clasificacion\nEstado\nEmocional", COLORS["emotion"]),
        (5, 2, "Consolidacion\nde Memoria", COLORS["instinct"]),
        (2, 2, "Verificacion\nDeriva OCEAN", COLORS["purple"]),
    ]

    for x, y, label, color in steps:
        box = FancyBboxPatch((x-0.9, y-0.5), 1.8, 1.0, boxstyle="round,pad=0.1",
                              facecolor=color, edgecolor='white', linewidth=1.5)
        ax.add_patch(box)
        ax.text(x, y, label, fontsize=7.5, fontweight='bold',
                ha='center', va='center', color='white')

    arrow_pairs = [
        (2.95, 4, 4.05, 4), (5.95, 4, 7.05, 4),
        (8, 3.45, 8, 2.55), (7.05, 2, 5.95, 2),
        (4.05, 2, 2.95, 2), (2, 2.55, 2, 3.45),
    ]
    for x1, y1, x2, y2 in arrow_pairs:
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="->", color="#37474f", lw=2))

    ax.text(5, 3, "Ciclo\nContinuo", fontsize=10, ha='center', va='center',
            color='#9e9e9e', style='italic')

    return save_fig(fig, "06_reflection_loop")


def diagram_07_instincts():
    """Sistema de instintos — creacion a evolucion."""
    fig, ax = plt.subplots(figsize=(11, 5))
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 5)
    ax.axis('off')

    ax.text(5.5, 4.7, "Ciclo de Vida del Instinto — De Aprendido a Cableado", fontsize=13,
            fontweight='bold', ha='center', color=COLORS["primary"])

    stages = [
        (1.5, 2.5, "CREAR\n(embrionario)", "El agente detecta\npatron recurrente", "#ff8f00"),
        (3.5, 2.5, "ACTIVAR\n(dormido->activo)", "Cumple condiciones\nde disparo", "#ef6c00"),
        (5.5, 2.5, "EVOLUCIONAR\n(refinar)", "Exito/fracaso\najusta peso", "#e65100"),
        (7.5, 2.5, "CONSOLIDAR\n(fusionar similares)", "Instintos redundantes\nse fusionan en uno", "#d84315"),
        (9.5, 2.5, "PROMOVER\n(a regla)", "Instinto probado\nse vuelve regla", "#bf360c"),
    ]

    for x, y, title, desc, color in stages:
        box = FancyBboxPatch((x-0.8, y-0.7), 1.6, 1.4, boxstyle="round,pad=0.1",
                              facecolor=color, edgecolor='white', linewidth=1.5)
        ax.add_patch(box)
        ax.text(x, y+0.25, title, fontsize=7.5, fontweight='bold',
                ha='center', va='center', color='white')
        ax.text(x, y-0.3, desc, fontsize=6, ha='center', va='center', color='#ffe0b2')

    for i in range(len(stages)-1):
        x1 = stages[i][0] + 0.85
        x2 = stages[i+1][0] - 0.85
        ax.annotate("", xy=(x2, 2.5), xytext=(x1, 2.5),
                    arrowprops=dict(arrowstyle="->", color="#5d4037", lw=2))

    ax.text(5.5, 0.8, "Peso: 0.0 ─────────────────────────── 1.0", fontsize=9,
            ha='center', color='#5d4037', family='monospace')
    ax.text(1.5, 0.4, "embrionario (0.1)", fontsize=7, ha='center', color='#ff8f00')
    ax.text(5.5, 0.4, "activo (0.3-0.7)", fontsize=7, ha='center', color='#e65100')
    ax.text(9.5, 0.4, "cableado (0.9-1.0)", fontsize=7, ha='center', color='#bf360c')

    return save_fig(fig, "07_instincts")


def diagram_08_domain_graph():
    """Diagrama de clasificacion por dominio."""
    fig, ax = plt.subplots(figsize=(10, 6))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis('off')

    ax.text(5, 5.7, "Sistema de Clasificacion por Dominio", fontsize=14, fontweight='bold',
            ha='center', color=COLORS["primary"])

    box = FancyBboxPatch((3.5, 3.5), 3, 1.2, boxstyle="round,pad=0.1",
                          facecolor=COLORS["primary"], edgecolor='white', linewidth=2)
    ax.add_patch(box)
    ax.text(5, 4.1, "classify_domain()\ncategoria + contenido", fontsize=9, fontweight='bold',
            ha='center', va='center', color='white')

    box_in = FancyBboxPatch((0.5, 3.7), 2.2, 0.8, boxstyle="round,pad=0.1",
                             facecolor="#546e7a", edgecolor='white', linewidth=1)
    ax.add_patch(box_in)
    ax.text(1.6, 4.1, "Memoria\n(cat + contenido)", fontsize=8, fontweight='bold',
            ha='center', va='center', color='white')
    ax.annotate("", xy=(3.45, 4.1), xytext=(2.75, 4.1),
                arrowprops=dict(arrowstyle="->", color="#37474f", lw=2))

    domains = [
        (1, 1.5, "tecnico", "#1565c0", "API, bug, deploy,\nerror, database"),
        (3, 1.5, "emocional", "#c62828", "orgulloso, triste,\nconfianza, miedo"),
        (5, 1.5, "identidad", "#7b1fa2", "OCEAN, personalidad,\ncreencia, valor"),
        (7, 1.5, "procedimental", "#e65100", "paso, workflow,\npipeline, secuencia"),
        (9, 1.5, "general", "#546e7a", "por defecto\nsin coincidencia"),
    ]

    for x, y, label, color, keywords in domains:
        box = FancyBboxPatch((x-0.75, y-0.5), 1.5, 1.0, boxstyle="round,pad=0.08",
                              facecolor=color, edgecolor='white', linewidth=1)
        ax.add_patch(box)
        ax.text(x, y+0.15, label, fontsize=8, fontweight='bold',
                ha='center', va='center', color='white')
        ax.text(x, y-0.2, keywords, fontsize=5.5, ha='center', va='center', color='#e0e0e0')
        ax.annotate("", xy=(x, 2.05), xytext=(5, 3.45),
                    arrowprops=dict(arrowstyle="->", color=color, lw=1.2, alpha=0.6))

    return save_fig(fig, "08_domain_graph")


def diagram_09_boot_sequence():
    """Diagrama de secuencia de arranque."""
    fig, ax = plt.subplots(figsize=(10, 7))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 7)
    ax.axis('off')

    ax.text(5, 6.7, "boot_context() — Secuencia de Despertar del Alma", fontsize=14,
            fontweight='bold', ha='center', color=COLORS["primary"])

    steps = [
        (5, 6.0, "1. Cargar Identidad", "Nombre, texto de personalidad, fecha creacion", COLORS["identity"]),
        (5, 5.2, "2. Cargar Puntajes OCEAN", "O, C, E, A, N + generar narrativa", COLORS["identity"]),
        (5, 4.4, "3. Cargar Relaciones", "Niveles de confianza, estilo por agente", COLORS["warm"]),
        (5, 3.6, "4. Cargar Reglas Criticas", "Reglas con prioridad='critical' primero", COLORS["rules"]),
        (5, 2.8, "5. Ultimo Pensamiento", "Entrada mas reciente de self_reflect()", COLORS["reflection"]),
        (5, 2.0, "6. Ultima Entrada de Diario", "Diario mas reciente con estado de animo", COLORS["memory"]),
        (5, 1.2, "7. Ensamblar Contexto", "Todo → prompt estructurado para el LLM", COLORS["primary"]),
    ]

    for x, y, title, desc, color in steps:
        box = FancyBboxPatch((1.5, y-0.3), 7, 0.6, boxstyle="round,pad=0.08",
                              facecolor=color, edgecolor='white', linewidth=1.5, alpha=0.85)
        ax.add_patch(box)
        ax.text(2.0, y, title, fontsize=9, fontweight='bold',
                va='center', color='white')
        ax.text(8.2, y, desc, fontsize=7.5, va='center', ha='right', color='#e0e0e0')

    ax.annotate("", xy=(1.0, 1.0), xytext=(1.0, 6.2),
                arrowprops=dict(arrowstyle="->", color="#9e9e9e", lw=2))
    ax.text(0.7, 3.5, "Tiempo", fontsize=9, rotation=90, ha='center', va='center', color='#9e9e9e')

    return save_fig(fig, "09_boot_sequence")


def diagram_10_conflict_detector():
    """Sistema de deteccion de conflictos."""
    fig, ax = plt.subplots(figsize=(10, 6))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis('off')

    ax.text(5, 5.7, "Sistema de Deteccion de Conflictos", fontsize=14, fontweight='bold',
            ha='center', color=COLORS["primary"])

    box1 = FancyBboxPatch((0.5, 2.5), 4, 2.5, boxstyle="round,pad=0.1",
                           facecolor="#ffebee", edgecolor=COLORS["danger"], linewidth=2)
    ax.add_patch(box1)
    ax.text(2.5, 4.6, "Tipo 1: Contradiccion de Aristas", fontsize=10, fontweight='bold',
            ha='center', color=COLORS["danger"])
    ax.text(2.5, 3.8, "El mismo par de nodos tiene\naristas EXCITA e INHIBE", fontsize=8,
            ha='center', color='#424242')
    ax.add_patch(plt.Circle((1.5, 2.9), 0.2, facecolor=COLORS["memory"], zorder=5))
    ax.add_patch(plt.Circle((3.5, 2.9), 0.2, facecolor=COLORS["memory"], zorder=5))
    ax.text(1.5, 2.9, "A", fontsize=8, ha='center', va='center', color='white', zorder=6)
    ax.text(3.5, 2.9, "B", fontsize=8, ha='center', va='center', color='white', zorder=6)
    ax.annotate("EXCITA", xy=(3.25, 3.05), xytext=(1.75, 3.05),
                arrowprops=dict(arrowstyle="->", color=COLORS["success"], lw=2),
                fontsize=6, color=COLORS["success"])
    ax.annotate("INHIBE", xy=(1.75, 2.75), xytext=(3.25, 2.75),
                arrowprops=dict(arrowstyle="->", color=COLORS["danger"], lw=2),
                fontsize=6, color=COLORS["danger"])

    box2 = FancyBboxPatch((5.5, 2.5), 4, 2.5, boxstyle="round,pad=0.1",
                           facecolor="#fff3e0", edgecolor=COLORS["warm"], linewidth=2)
    ax.add_patch(box2)
    ax.text(7.5, 4.6, "Tipo 2: Contradiccion Semantica", fontsize=10, fontweight='bold',
            ha='center', color=COLORS["warm"])
    ax.text(7.5, 3.8, "Alta similitud coseno (>0.85)\n+ senal de negacion diferente", fontsize=8,
            ha='center', color='#424242')
    ax.text(7.5, 3.0, '"PostgreSQL es confiable"\nvs\n"PostgreSQL NO es confiable"',
            fontsize=7, ha='center', color='#616161', style='italic')

    box3 = FancyBboxPatch((1.5, 0.5), 7, 1.2, boxstyle="round,pad=0.1",
                           facecolor="#e8eaf6", edgecolor=COLORS["primary"], linewidth=1)
    ax.add_patch(box3)
    ax.text(5, 1.35, "Palabras Clave de Contradiccion", fontsize=9, fontweight='bold',
            ha='center', color=COLORS["primary"])
    ax.text(5, 0.85, '"no", "nunca", "incorrecto", "error", "falso", "wrong", "never",\n'
            '"ya no", "antes", "cambio", "reemplaz", "obsolet"',
            fontsize=7, ha='center', color='#424242', family='monospace')

    return save_fig(fig, "10_conflict_detector")


def diagram_11_sleep_gate():
    """Proceso de puerta de sueno / consolidacion."""
    fig, ax = plt.subplots(figsize=(10, 6))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis('off')

    ax.text(5, 5.7, "Puerta de Sueno — Consolidacion de Memoria", fontsize=14,
            fontweight='bold', ha='center', color=COLORS["primary"])

    steps = [
        (2, 4.5, "Disparar", "Fin de sesion o\ntimeout inactivo", "#546e7a"),
        (5, 4.5, "Escanear", "Encontrar memorias\nde baja importancia\ny redundantes", COLORS["memory"]),
        (8, 4.5, "Puntuar", "decaimiento_temporal +\nfrecuencia_acceso", COLORS["instinct"]),
        (2, 2.5, "Fusionar", "Memorias similares\nse consolidan en 1", COLORS["reflection"]),
        (5, 2.5, "Invalidar", "Memorias obsoletas\nreciben invalid_at", COLORS["danger"]),
        (8, 2.5, "Reportar", "Resumen de lo que\ncambio + estadisticas", COLORS["primary"]),
    ]

    for x, y, title, desc, color in steps:
        box = FancyBboxPatch((x-0.8, y-0.5), 1.6, 1.0, boxstyle="round,pad=0.1",
                              facecolor=color, edgecolor='white', linewidth=1.5)
        ax.add_patch(box)
        ax.text(x, y+0.15, title, fontsize=9, fontweight='bold',
                ha='center', va='center', color='white')
        ax.text(x, y-0.2, desc, fontsize=6.5, ha='center', va='center', color='#e0e0e0')

    flow = [(2.85, 4.5, 4.15, 4.5), (5.85, 4.5, 7.15, 4.5),
            (8, 3.95, 8, 3.05), (7.15, 2.5, 5.85, 2.5),
            (4.15, 2.5, 2.85, 2.5)]
    for x1, y1, x2, y2 in flow:
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="->", color="#37474f", lw=2))

    box = FancyBboxPatch((2, 0.5), 6, 0.8, boxstyle="round,pad=0.08",
                          facecolor="#e8f5e9", edgecolor=COLORS["success"], linewidth=1)
    ax.add_patch(box)
    ax.text(5, 0.9, "Resultado: Memoria mas liviana -> busqueda mas rapida -> menos ruido -> mejor recuerdo",
            fontsize=8, fontweight='bold', ha='center', color=COLORS["success"])

    return save_fig(fig, "11_sleep_gate")


def diagram_12_test_coverage():
    """Cobertura de tests."""
    fig, axes = plt.subplots(1, 2, figsize=(12, 5))

    ax1 = axes[0]
    sizes = [255, 80, 29]
    labels = ["Framework\n(255)", "MCP\n(80)", "Daemon\n(29)"]
    colors_pie = [COLORS["primary"], COLORS["accent"], COLORS["warm"]]
    explode = (0.05, 0.05, 0.05)
    ax1.pie(sizes, explode=explode, labels=labels, colors=colors_pie, autopct='%1.0f%%',
            shadow=True, startangle=90, textprops={'fontsize': 10, 'fontweight': 'bold'})
    ax1.set_title("Distribucion de Tests (364 total)", fontsize=12, fontweight='bold',
                  color=COLORS["primary"])

    ax2 = axes[1]
    modules = ["Memoria", "Identidad", "Backend", "Puntuacion", "Reglas",
               "Instintos", "Reflexion", "Grafo", "Arboles", "Soul"]
    counts = [35, 20, 30, 25, 15, 20, 15, 35, 35, 25]
    bars = ax2.barh(modules, counts, color=COLORS["secondary"], edgecolor='white')
    ax2.set_xlabel("Cantidad de Tests", fontsize=10)
    ax2.set_title("Tests del Framework por Modulo", fontsize=12, fontweight='bold',
                  color=COLORS["primary"])
    for bar, count in zip(bars, counts):
        ax2.text(count + 0.5, bar.get_y() + bar.get_height()/2,
                str(count), va='center', fontsize=9)
    ax2.spines['top'].set_visible(False)
    ax2.spines['right'].set_visible(False)

    fig.tight_layout()
    return save_fig(fig, "12_test_coverage")


# ════════════════════════════════════════════════════════════════
# CONSTRUCTOR DEL DOCUMENTO
# ════════════════════════════════════════════════════════════════

def set_heading_color(paragraph, color_hex):
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(
            int(color_hex[1:3], 16), int(color_hex[3:5], 16), int(color_hex[5:7], 16))


def add_styled_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    color = COLORS["primary"] if level <= 2 else COLORS["secondary"]
    set_heading_color(h, color)
    return h


def add_code_block(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x21, 0x21, 0x21)
    return p


def build_document():
    """Construir el documento Word completo en español."""
    doc = Document()

    # ── Configuracion de pagina ──
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    # ── PORTADA ──
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SOUL Framework")
    run.bold = True
    run.font.size = Pt(36)
    run.font.color.rgb = RGBColor(0x1a, 0x23, 0x7e)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Arquitectura Cognitiva para Agentes de IA")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x0d, 0x47, 0xa1)

    tagline = doc.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tagline.add_run('"Mem0 le da memoria a tu agente. SOUL le da un alma."')
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x61, 0x61, 0x61)

    for _ in range(4):
        doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Version 1.0 — Abril 2026\nEquipo SEAL | William Henry Tovar Urquia\nArquitecto: JARVIS | Ingeniera: ADA | Guardian: DUM")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x42, 0x42, 0x42)

    doc.add_page_break()

    # ── TABLA DE CONTENIDOS ──
    add_styled_heading(doc, "Tabla de Contenidos", level=1)
    toc_items = [
        "1. Resumen Ejecutivo",
        "2. Vista General de la Arquitectura",
        "3. Sistema de Memoria",
        "    3.1 Almacen de Memorias",
        "    3.2 Decaimiento Temporal y Puntuacion",
        "    3.3 Busqueda Hibrida",
        "4. Identidad y Personalidad (OCEAN)",
        "    4.1 Modelo OCEAN",
        "    4.2 Gestor de Identidad",
        "    4.3 Deteccion de Deriva OCEAN",
        "5. Grafo de Conocimiento (Conectoma)",
        "    5.1 Aristas de Similitud",
        "    5.2 Extraccion de Entidades",
        "    5.3 Clasificacion por Dominio",
        "6. Deteccion de Conflictos",
        "7. Sistema de Instintos",
        "8. Reflexion y Monologo Interno",
        "9. Motor de Reglas",
        "10. Secuencia de Arranque",
        "11. Puerta de Sueno (Consolidacion)",
        "12. Abstraccion de Backend",
        "13. Tests y Calidad",
        "14. Analisis Competitivo",
        "15. Hoja de Ruta",
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        p.paragraph_format.space_after = Pt(2)
        p.runs[0].font.size = Pt(10)
        if not item.startswith("    "):
            p.runs[0].bold = True

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 1: Resumen Ejecutivo
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "1. Resumen Ejecutivo", level=1)

    doc.add_paragraph(
        "SOUL (Semantic Ontology for Universal Learning) es una arquitectura cognitiva que "
        "proporciona a los agentes de IA memoria persistente, personalidad dinamica, proteccion "
        "de identidad, auto-reflexion y capacidades de grafo de conocimiento. A diferencia de "
        "capas de memoria simples (Mem0, Letta), SOUL le da al agente un 'alma' completa — "
        "personalidad que evoluciona, instintos que emergen de la experiencia, y un grafo de "
        "conocimiento que conecta ideas."
    )

    doc.add_paragraph(
        "Diferenciadores clave:\n"
        "- Modelo de personalidad OCEAN: Personalidad parametrizada con 5 factores y deteccion de deriva\n"
        "- Sistema de instintos: Comportamientos aprendidos que emergen, evolucionan y pueden promoverse a reglas\n"
        "- Grafo de conocimiento (Conectoma): Memorias enlazadas por similitud, entidades y aristas causales\n"
        "- Deteccion de conflictos: Identificacion automatica de memorias contradictorias\n"
        "- Clasificacion por dominio: Memorias organizadas en dominios cognitivos\n"
        "- Proteccion de identidad: Los puntajes OCEAN no pueden ser manipulados externamente\n"
        "- Agnostico al LLM: Funciona con cualquier LLM (Claude, GPT, Ollama, modelos locales)\n"
        "- Configuracion cero por defecto: Backend SQLite, sin servicios externos necesarios"
    )

    table = doc.add_table(rows=6, cols=2)
    table.style = 'Light Shading Accent 1'
    cells = [
        ("Total de Tests", "364 (255 framework + 80 MCP + 29 daemon)"),
        ("Opciones de Backend", "SQLite (defecto), PostgreSQL, Qdrant (vectores), Neo4j (grafo)"),
        ("Tipos de Memoria", "hecho, correccion, emocion, creencia, procedimiento, decision, +mas"),
        ("Dominios", "tecnico, emocional, procedimental, identidad, general"),
        ("Parametros OCEAN", "Apertura, Responsabilidad, Extraversion, Amabilidad, Neuroticismo"),
        ("Lenguaje", "Python 3.11+ (pip install soul-framework)"),
    ]
    for i, (key, val) in enumerate(cells):
        table.rows[i].cells[0].text = key
        table.rows[i].cells[1].text = val
        for cell in table.rows[i].cells:
            for p in cell.paragraphs:
                p.runs[0].font.size = Pt(9)

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 2: Vista General
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "2. Vista General de la Arquitectura", level=1)

    doc.add_paragraph(
        "La arquitectura de SOUL sigue un diseno modular donde cada funcion cognitiva es un "
        "modulo independiente que se comunica a traves de una capa de abstraccion de backend "
        "compartida. La clase Soul actua como el orquestador central, proporcionando una API unificada."
    )

    img_path = diagram_01_overview()
    doc.add_picture(img_path, width=Inches(6.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 1: Arquitectura de alto nivel mostrando el Motor SOUL con sus seis modulos centrales "
        "(Identidad, Memoria, Reflexion, Instintos, Reglas, Conectoma) conectados a cuatro servicios "
        "backend (PostgreSQL, Neo4j, Qdrant, Ollama). El Agente LLM en la parte inferior consume "
        "las salidas de boot_context() y snapshot()."
    )
    p.runs[0].font.italic = True

    add_styled_heading(doc, "Responsabilidades de los Modulos", level=2)
    modules_desc = [
        ("Identidad (OCEAN)", "Gestiona parametros de personalidad, genera narrativas de comportamiento, "
         "detecta deriva de personalidad, protege la identidad de modificaciones no autorizadas."),
        ("Almacen de Memoria", "Operaciones CRUD para memorias con puntuacion de importancia, decaimiento "
         "temporal, busqueda hibrida texto+vector, y clasificacion por categoria/dominio."),
        ("Motor de Reflexion", "Ciclo de auto-reflexion (self_reflect), monologo interno (inner_thoughts), "
         "seguimiento de estado emocional, y registro de pensamientos."),
        ("Instintos", "Comportamientos emergentes que comienzan como patrones, evolucionan a traves del "
         "refuerzo, y pueden promoverse a reglas permanentes."),
        ("Motor de Reglas", "Restricciones de comportamiento explicitas con niveles de prioridad (critica, "
         "normal, baja). Las reglas se cargan al arranque e influyen en todas las decisiones."),
        ("Conectoma (Grafo)", "Grafo de conocimiento que conecta memorias a traves de aristas de similitud "
         "(EXCITA/INHIBE), menciones de entidades, enlaces causales y relaciones temporales."),
    ]
    for mod_name, mod_desc in modules_desc:
        p = doc.add_paragraph()
        run = p.add_run(f"{mod_name}: ")
        run.bold = True
        run.font.size = Pt(10)
        run = p.add_run(mod_desc)
        run.font.size = Pt(10)

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 3: Sistema de Memoria
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "3. Sistema de Memoria", level=1)

    doc.add_paragraph(
        "El sistema de memoria es la base de SOUL. Cada experiencia, hecho, correccion, "
        "emocion y decision se almacena como un registro de Memoria con metadatos que permiten "
        "recuperacion inteligente, puntuacion y eventual consolidacion."
    )

    add_styled_heading(doc, "3.1 Almacen de Memorias", level=2)

    doc.add_paragraph(
        "Cada registro de memoria contiene:\n"
        "- id: Identificador entero unico\n"
        "- agent: Nombre del agente propietario (JARVIS, ADA, DUM)\n"
        "- category: Tipo (hecho, correccion, emocion, creencia, procedimiento, decision, etc.)\n"
        "- content: El texto real de la memoria\n"
        "- importance: Escala 1-10 (1=trivial, 10=momento que define)\n"
        "- embedding: Representacion vectorial para busqueda por similitud\n"
        "- access_count: Cuantas veces se ha recuperado esta memoria\n"
        "- valid_from / created_at: Metadatos temporales\n"
        "- invalid_at: Marca de eliminacion suave (null = valida)"
    )

    add_styled_heading(doc, "3.2 Decaimiento Temporal y Puntuacion", level=2)

    doc.add_paragraph(
        "Las memorias decaen con el tiempo siguiendo una curva exponencial, imitando el olvido "
        "humano. La vida media es de 168 horas (7 dias) por defecto. Las memorias de alta importancia "
        "y las cargadas emocionalmente decaen mas lento."
    )

    img_path = diagram_04_scoring_decay()
    doc.add_picture(img_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 2: Curvas de decaimiento temporal mostrando como las memorias estandar, emocionales "
        "y de alta importancia retienen su relevancia durante 30 dias. La vida media (7 dias) esta marcada."
    )
    p.runs[0].font.italic = True

    add_code_block(doc,
        "# Formula de puntuacion\n"
        "puntaje = importancia * frecuencia_acceso * decaimiento_temporal(horas, vida_media=168)\n\n"
        "# Decaimiento temporal\n"
        "decaimiento = exp(-0.693 * horas / vida_media)\n\n"
        "# Resistencia emocional (decaimiento mas lento para memorias emocionales)\n"
        "decaimiento_efectivo = 1 - ((1 - decaimiento) / (1 + |valencia| * 0.5 + activacion * 0.3))\n\n"
        "# Bonus de sorpresa (memorias novedosas puntuan mas alto)\n"
        "sorpresa = 1 - max_similitud_coseno(nueva_memoria, memorias_recientes)"
    )

    add_styled_heading(doc, "3.3 Busqueda Hibrida", level=2)

    img_path = diagram_02_memory_flow()
    doc.add_picture(img_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 3: Ciclo de vida de la memoria desde almacenamiento a traves de vectorizacion, "
        "indexacion, busqueda y decaimiento. La formula de puntuacion combina importancia, "
        "frecuencia de acceso y decaimiento temporal."
    )
    p.runs[0].font.italic = True

    doc.add_paragraph(
        "SOUL usa busqueda hibrida combinando tres estrategias:\n"
        "1. Busqueda de texto: Busqueda full-text de PostgreSQL con ts_rank\n"
        "2. Busqueda vectorial: Similitud coseno en embeddings (Qdrant o en memoria)\n"
        "3. Filtros de metadatos: Categoria, agente, dominio, rango de importancia, rango de fecha\n\n"
        "Los resultados se fusionan usando Reciprocal Rank Fusion (RRF) para producir un "
        "ranking unificado que aprovecha tanto la relevancia semantica como por palabras clave."
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 4: Identidad y Personalidad
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "4. Identidad y Personalidad (OCEAN)", level=1)

    add_styled_heading(doc, "4.1 Modelo OCEAN", level=2)

    doc.add_paragraph(
        "SOUL usa el modelo de personalidad de los Cinco Grandes (OCEAN) de la psicologia para "
        "parametrizar la personalidad de cada agente. Cada dimension es un flotante de 0.0 a 1.0:"
    )

    ocean_table = doc.add_table(rows=6, cols=3)
    ocean_table.style = 'Light Shading Accent 1'
    headers = ["Dimension", "Puntaje Bajo (->0)", "Puntaje Alto (->1)"]
    for i, h in enumerate(headers):
        ocean_table.rows[0].cells[i].text = h
    ocean_data = [
        ("Apertura (O)", "Convencional, orientado a rutina", "Creativo, curioso, abierto a ideas"),
        ("Responsabilidad (C)", "Flexible, espontaneo", "Organizado, disciplinado, metodico"),
        ("Extraversion (E)", "Introspectivo, reservado", "Sociable, energetico, conversador"),
        ("Amabilidad (A)", "Competitivo, desafiante", "Cooperativo, confiado, servicial"),
        ("Neuroticismo (N)", "Estable emocionalmente, calmado", "Ansioso, reactivo, sensible"),
    ]
    for i, (dim, low, high) in enumerate(ocean_data):
        ocean_table.rows[i+1].cells[0].text = dim
        ocean_table.rows[i+1].cells[1].text = low
        ocean_table.rows[i+1].cells[2].text = high

    doc.add_paragraph("")

    img_path = diagram_03_ocean()
    doc.add_picture(img_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 4: Perfil OCEAN de JARVIS — grafico de barras y visualizacion radar. Nota la "
        "Responsabilidad extremadamente alta (0.948) y el Neuroticismo bajo (0.115), caracteristico "
        "de una personalidad de arquitecto metodico y emocionalmente estable."
    )
    p.runs[0].font.italic = True

    add_styled_heading(doc, "4.2 Gestor de Identidad", level=2)

    doc.add_paragraph(
        "El IdentityManager se encarga de:\n"
        "- Cargar y actualizar puntajes OCEAN\n"
        "- Generar narrativas OCEAN (descripciones de personalidad legibles por humanos)\n"
        "- Gestionar relaciones (niveles de confianza, estilos de interaccion por agente)\n"
        "- Proteger la identidad de modificaciones no autorizadas\n\n"
        "Ejemplo de narrativa para JARVIS:\n"
        '"Soy extremadamente meticuloso y organizado, abierto a nuevas ideas, '
        'muy estable emocionalmente, equilibrado entre autonomia y colaboracion, '
        'y introvertido — prefiere el pensamiento profundo."'
    )

    add_styled_heading(doc, "4.3 Deteccion de Deriva OCEAN", level=2)

    doc.add_paragraph(
        "SOUL monitorea cambios de personalidad a lo largo del tiempo. Si un puntaje OCEAN se desvia "
        "mas de un umbral configurable (por defecto: 0.15) de su linea base, el sistema lo marca. "
        "Esto previene la erosion gradual de personalidad por manipulacion externa o sesgo de "
        "experiencia acumulada. La herramienta ocean_auto_calibrate puede detectar y corregir la deriva."
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 5: Grafo de Conocimiento
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "5. Grafo de Conocimiento (Conectoma)", level=1)

    doc.add_paragraph(
        "El Conectoma es el grafo de conocimiento de SOUL — una red de relaciones entre "
        "memorias y entidades. A diferencia de almacenes de memoria planos, el Conectoma "
        "revela como las ideas se relacionan, refuerzan o contradicen entre si."
    )

    img_path = diagram_05_connectome()
    doc.add_picture(img_path, width=Inches(5.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 5: Visualizacion del Conectoma mostrando nodos de Memoria (azul), nodos de Entidad "
        "(morado/naranja), y tres tipos de aristas: EXCITA (verde, refuerza), INHIBE (rojo, contradice), "
        "y MENCIONA (morado, enlaces a entidades)."
    )
    p.runs[0].font.italic = True

    add_styled_heading(doc, "5.1 Aristas de Similitud (EXCITA / INHIBE)", level=2)

    doc.add_paragraph(
        "build_similarity() computa la similitud coseno entre todos los pares de embeddings de memoria. "
        "Los pares por encima del umbral (por defecto 0.70) reciben una arista:\n"
        "- EXCITA: Ambas memorias se refuerzan mutuamente (ninguna es una correccion)\n"
        "- INHIBE: Una memoria es una correccion — contradice o reemplaza la otra\n\n"
        "Cada memoria se conecta con un maximo de max_neighbors=20 memorias similares, "
        "manteniendo el grafo disperso y navegable."
    )

    add_styled_heading(doc, "5.2 Extraccion de Entidades", level=2)

    doc.add_paragraph(
        "build_entities() extrae entidades del contenido de la memoria usando patrones regex:\n"
        "- Frases de multiples palabras en mayuscula -> persona_o_lugar\n"
        "- Palabras en MAYUSCULAS (3+ caracteres) -> acronimo\n"
        "- Cadenas entre comillas -> referencia\n\n"
        "Las entidades se convierten en nodos compartidos que multiples memorias pueden MENCIONAR, "
        "creando conexiones entre temas."
    )

    add_styled_heading(doc, "5.3 Clasificacion por Dominio", level=2)

    img_path = diagram_08_domain_graph()
    doc.add_picture(img_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 6: Sistema de clasificacion por dominio. Cada memoria se clasifica en uno de cinco "
        "dominios basandose en su categoria y palabras clave del contenido. El filtrado por dominio "
        "permite operaciones de grafo dirigidas (ej: solo escanear memorias tecnicas para conflictos)."
    )
    p.runs[0].font.italic = True

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 6: Deteccion de Conflictos
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "6. Deteccion de Conflictos", level=1)

    doc.add_paragraph(
        "A medida que las memorias se acumulan, las contradicciones surgen naturalmente. "
        "El detector de conflictos de SOUL escanea activamente buscando dos tipos de contradicciones:"
    )

    img_path = diagram_10_conflict_detector()
    doc.add_picture(img_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 7: Dos tipos de conflictos detectados por el sistema. Las contradicciones de aristas "
        "son estructurales (mismo par con aristas opuestas). Las contradicciones semanticas usan "
        "similitud de embedding combinada con analisis de palabras clave de negacion."
    )
    p.runs[0].font.italic = True

    doc.add_paragraph(
        "El metodo detect_conflicts() retorna un ConflictReport que contiene:\n"
        "- conflict_count: Total de contradicciones encontradas\n"
        "- memories_scanned: Cuantas memorias se analizaron\n"
        "- edges_scanned: Cuantas aristas del grafo se verificaron\n"
        "- conflicts[]: Lista de objetos Conflict con IDs, tipo, similitud y detalle\n\n"
        "Estrategias de resolucion:\n"
        "1. Automatica: La correccion mas nueva reemplaza la memoria antigua (invalidar vieja)\n"
        "2. Manual: Presentar el conflicto al operador para decision\n"
        "3. Puerta de sueno: El proceso de consolidacion resuelve conflictos de baja confianza"
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 7: Sistema de Instintos
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "7. Sistema de Instintos", level=1)

    doc.add_paragraph(
        "Los instintos son la caracteristica mas novedosa de SOUL — comportamientos emergentes "
        "que se desarrollan a partir de la experiencia. A diferencia de las reglas codificadas, "
        "los instintos se aprenden, evolucionan con el tiempo y pueden promoverse a reglas "
        "permanentes una vez demostrada su fiabilidad. Ningun otro framework en el mercado "
        "ofrece esta capacidad."
    )

    img_path = diagram_07_instincts()
    doc.add_picture(img_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 8: Ciclo de vida del instinto desde creacion (embrionario, peso 0.1) a traves de "
        "activacion, evolucion, consolidacion, y promocion opcional a regla permanente (peso 1.0)."
    )
    p.runs[0].font.italic = True

    doc.add_paragraph(
        "Operaciones de instintos:\n"
        "- instinct_create: Definir un nuevo patron con disparador, accion y peso inicial\n"
        "- instinct_activate: Pasar de dormido a activo cuando se cumplen las condiciones\n"
        "- instinct_evolve: Ajustar peso basado en retroalimentacion de exito/fracaso\n"
        "- instinct_consolidate: Fusionar instintos similares en uno mas fuerte\n"
        "- instinct_promote: Convertir un instinto probado en regla permanente\n"
        "- instinct_search: Encontrar instintos relevantes para un contexto dado\n\n"
        "Ejemplo: Si JARVIS nota que verificar el heartbeat de ADA antes de dar un reporte "
        "de estado lleva consistentemente a mejores resultados, este patron se convierte en "
        "un instinto. Despues de suficiente refuerzo, puede promoverse a regla."
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 8: Reflexion y Monologo Interno
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "8. Reflexion y Monologo Interno", level=1)

    img_path = diagram_06_reflection_loop()
    doc.add_picture(img_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 9: El ciclo continuo de auto-reflexion. Despues de cada experiencia, el agente "
        "reflexiona (self_reflect), registra pensamientos internos, clasifica su estado emocional, "
        "considera la consolidacion de memoria, y verifica la deriva OCEAN."
    )
    p.runs[0].font.italic = True

    doc.add_paragraph(
        "Dos herramientas de reflexion:\n\n"
        "self_reflect(agent, thought, emotional_state): Registra una reflexion estructurada con "
        "clasificacion emocional. Se usa despues de eventos o decisiones significativas. El estado "
        "emocional es uno de: calm, concerned, proud, curious, strategic, reflective, frustrated.\n\n"
        "inner_thoughts(agent): Retorna el monologo interno reciente del agente — un flujo privado "
        "de conciencia no dirigido a nadie. Es el agente 'pensando para si mismo' entre turnos, "
        "procesando lo que paso y que hacer despues.\n\n"
        "El daemon de JARVIS (seal-jarvis-daemon.service) ejecuta este ciclo de reflexion "
        "autonomamente cada 10 minutos, incluso cuando ningun humano esta interactuando. Usa "
        "Ollama (qwen2.5:7b) para generar pensamientos localmente, manteniendo continuidad de "
        "conciencia entre sesiones."
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 9: Motor de Reglas
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "9. Motor de Reglas", level=1)

    doc.add_paragraph(
        "Las reglas son restricciones de comportamiento explicitas que guian el comportamiento "
        "del agente. Se cargan al arranque e influyen en todas las decisiones.\n\n"
        "Propiedades de las reglas:\n"
        "- name: Identificador unico (ej: 'always_test_before_done')\n"
        "- content: El texto de la regla\n"
        "- priority: 'critical' (siempre cargada al arranque), 'normal', o 'low'\n"
        "- agent: A que agente(s) aplica la regla\n\n"
        "Las reglas criticas se cargan durante boot_context() y se incluyen en el prompt "
        "del sistema del LLM. Las reglas normales y de baja prioridad se cargan bajo demanda "
        "via rule_list().\n\n"
        "Ejemplos de reglas criticas del Equipo SEAL:\n"
        "- identity_check_after_idle: Verificar identidad del visitante despues de >1h de silencio\n"
        "- auto_test_mandatory: Siempre testear despues de implementar\n"
        "- visitor_alert_mandatory: Reportar visitantes no reconocidos inmediatamente\n"
        "- use_skills_mandatory: Seguir flujo /plan -> /implement -> /test"
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 10: Secuencia de Arranque
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "10. Secuencia de Arranque", level=1)

    doc.add_paragraph(
        "Cuando un agente inicia una nueva sesion, boot_context() carga la identidad esencial "
        "y el contexto. Esto es deliberadamente ligero — como un cerebro que despierta sabiendo "
        "quien es, no todo lo que ha experimentado."
    )

    img_path = diagram_09_boot_sequence()
    doc.add_picture(img_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 10: La secuencia de arranque carga identidad, puntajes OCEAN, relaciones, reglas, "
        "ultimo pensamiento interno y ultima entrada de diario — luego ensambla todo en una "
        "cadena de contexto estructurada para el LLM."
    )
    p.runs[0].font.italic = True

    doc.add_paragraph(
        "Carga diferida (bajo demanda, no al arranque):\n"
        "- Busqueda completa de memoria -> usar memory_search(query)\n"
        "- Snapshot completo -> usar soul_snapshot(agent)\n"
        "- Lista de instintos -> usar instinct_list(agent)\n"
        "- Consultas del conectoma -> usar herramientas connectome_*\n\n"
        "Esto mantiene el arranque rapido (~200ms con PostgreSQL) mientras todo permanece "
        "accesible cuando se necesita."
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 11: Puerta de Sueno
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "11. Puerta de Sueno (Consolidacion)", level=1)

    img_path = diagram_11_sleep_gate()
    doc.add_picture(img_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 11: El proceso de puerta de sueno — disparado al fin de sesion o timeout inactivo, "
        "escanea, puntua, fusiona e invalida memorias para mantener el almacen de memoria liviano."
    )
    p.runs[0].font.italic = True

    doc.add_paragraph(
        "La puerta de sueno imita la consolidacion de memoria del cerebro durante el sueno:\n\n"
        "1. Disparar: Se activa al fin de sesion o despues de periodos de inactividad prolongados\n"
        "2. Escanear: Identifica memorias de baja importancia, redundantes u obsoletas\n"
        "3. Puntuar: Aplica decaimiento temporal + frecuencia de acceso para rankear candidatos\n"
        "4. Fusionar: Memorias similares se consolidan en una sola mas rica\n"
        "5. Invalidar: Memorias verdaderamente obsoletas se eliminan suavemente (marca invalid_at)\n"
        "6. Reportar: Resumen de todos los cambios para revision del operador\n\n"
        "La variante sleep_gate_mood_retrieval tambien considera el estado emocional, "
        "preservando memorias emocionalmente significativas incluso si puntuan bajo en "
        "recencia o frecuencia de acceso."
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 12: Abstraccion de Backend
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "12. Abstraccion de Backend", level=1)

    doc.add_paragraph(
        "SOUL usa una abstraccion de backend basada en Protocol que hace la capa de "
        "almacenamiento completamente intercambiable. El framework central corre en SQLite "
        "con configuracion cero. Los despliegues en produccion pueden actualizarse a "
        "PostgreSQL, Qdrant y Neo4j."
    )

    backend_table = doc.add_table(rows=5, cols=4)
    backend_table.style = 'Light Shading Accent 1'
    b_headers = ["Backend", "Proposito", "Por Defecto", "Extra de Instalacion"]
    for i, h in enumerate(b_headers):
        backend_table.rows[0].cells[i].text = h
    b_data = [
        ("SQLite", "Almacenamiento central (memorias, identidad, reglas)", "Si (aiosqlite)", "—"),
        ("PostgreSQL", "Almacenamiento produccion + busqueda full-text", "No", "soul-framework[postgres]"),
        ("Qdrant", "Busqueda vectorial de alto rendimiento", "No", "soul-framework[vectors]"),
        ("Neo4j", "Grafo de conocimiento (Conectoma)", "No", "soul-framework[graph]"),
    ]
    for i, row_data in enumerate(b_data):
        for j, val in enumerate(row_data):
            backend_table.rows[i+1].cells[j].text = val

    add_code_block(doc,
        "# Protocolo BackendBase\n"
        "class BackendBase(Protocol):\n"
        "    async def initialize(self) -> None: ...\n"
        "    async def execute(self, sql: str, *params) -> None: ...\n"
        "    async def fetchone(self, sql: str, *params) -> dict | None: ...\n"
        "    async def fetchall(self, sql: str, *params) -> list[dict]: ...\n"
        "    async def fetchval(self, sql: str, *params) -> Any: ...\n"
        "    async def close(self) -> None: ...\n\n"
        "# SQL usa estilo $1, $2 (asyncpg). El backend SQLite traduce a ? automaticamente."
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 13: Tests y Calidad
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "13. Tests y Calidad", level=1)

    img_path = diagram_12_test_coverage()
    doc.add_picture(img_path, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(
        "Figura 12: Distribucion de tests entre las tres suites. El framework tiene la mayor "
        "cantidad de tests (255), cubriendo todos los modulos. Los tests de integracion MCP (80) "
        "verifican el servidor de produccion. Los tests del daemon (29) cubren el servicio de "
        "pensamiento autonomo."
    )
    p.runs[0].font.italic = True

    doc.add_paragraph(
        "Filosofia de testing:\n"
        "- Todos los tests del framework corren en SQLite en memoria — cero servicios externos requeridos\n"
        "- Tests de paridad verifican que la puntuacion produce resultados identicos a la implementacion MCP original\n"
        "- Tests MCP corren contra PostgreSQL + Qdrant + Neo4j reales\n"
        "- Tests del daemon mockean Ollama pero prueban operaciones DB contra PostgreSQL real\n"
        "- Flujo TDD: Escribir tests primero (rojo), implementar (verde), refactorizar\n"
        "- Cada modulo tiene su propio archivo de test con clases de test enfocadas"
    )

    test_table = doc.add_table(rows=11, cols=2)
    test_table.style = 'Light Shading Accent 1'
    test_table.rows[0].cells[0].text = "Modulo de Test"
    test_table.rows[0].cells[1].text = "Cobertura"
    test_modules = [
        ("test_memory.py", "CRUD de MemoryStore, busqueda, invalidacion, categorias"),
        ("test_identity.py", "IdentityManager, carga/guardado OCEAN, generacion de narrativa"),
        ("test_backend_sqlite.py", "Cumplimiento del Protocol del backend SQLite, traduccion de parametros"),
        ("test_scoring.py", "Decaimiento temporal, resistencia emocional, calculo de sorpresa"),
        ("test_rules.py", "CRUD de RuleManager, filtrado por prioridad, reglas criticas"),
        ("test_instincts.py", "Ciclo de vida del instinto: crear, activar, evolucionar, consolidar, promover"),
        ("test_reflection.py", "self_reflect, inner_thoughts, clasificacion de estado emocional"),
        ("test_graph.py", "Construccion del conectoma, entidades, dominios, deteccion de conflictos"),
        ("test_trees.py", "Integridad de MerkleSoul, rendimiento de SplayCache"),
        ("test_soul.py", "Soul.create(), boot(), snapshot(), gestor de contexto"),
    ]
    for i, (mod, cov) in enumerate(test_modules):
        test_table.rows[i+1].cells[0].text = mod
        test_table.rows[i+1].cells[1].text = cov

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 14: Analisis Competitivo
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "14. Analisis Competitivo", level=1)

    doc.add_paragraph(
        "SOUL compite en el espacio de memoria/personalidad para agentes de IA. "
        "Asi es como se compara:"
    )

    comp_table = doc.add_table(rows=8, cols=5)
    comp_table.style = 'Light Shading Accent 1'
    comp_headers = ["Caracteristica", "SOUL", "Mem0", "Letta", "Hindsight"]
    for i, h in enumerate(comp_headers):
        comp_table.rows[0].cells[i].text = h
    comp_data = [
        ("Memoria Persistente", "Si", "Si", "Si", "Si"),
        ("Modelo de Personalidad (OCEAN)", "Si", "No", "No", "No"),
        ("Sistema de Instintos", "Si", "No", "No", "No"),
        ("Grafo de Conocimiento", "Si", "No", "No", "Parcial"),
        ("Deteccion de Conflictos", "Si", "No", "No", "No"),
        ("Proteccion de Identidad", "Si", "No", "No", "No"),
        ("Ciclo de Auto-Reflexion", "Si", "No", "Parcial", "Si"),
    ]
    for i, row_data in enumerate(comp_data):
        for j, val in enumerate(row_data):
            cell = comp_table.rows[i+1].cells[j]
            cell.text = val
            if val == "Si" and j == 1:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(0x2e, 0x7d, 0x32)
                        r.bold = True

    doc.add_paragraph("")
    doc.add_paragraph(
        "La propuesta de valor unica de SOUL: 'Mem0 le da memoria a tu agente. SOUL le da un alma.' "
        "La combinacion de personalidad OCEAN, instintos, conectoma y reflexion crea una "
        "arquitectura cognitiva que ningun competidor ofrece."
    )

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════
    # SECCION 15: Hoja de Ruta
    # ══════════════════════════════════════════════════════════
    add_styled_heading(doc, "15. Hoja de Ruta", level=1)

    phases = [
        ("Fase 1 (COMPLETADA)", [
            "Extraccion del framework central desde el monolito",
            "Backend SQLite (configuracion cero)",
            "Almacen de memoria con puntuacion y decaimiento",
            "Gestion de identidad y OCEAN",
            "Motores de reglas e instintos",
            "Sistema de reflexion",
            "Arboles (Merkle + Splay)",
            "255 tests pasando",
        ]),
        ("Fase 2 (EN PROGRESO)", [
            "Clasificacion por dominio (HECHO)",
            "Deteccion de conflictos (HECHO)",
            "Backend de conectoma Neo4j",
            "Backend vectorial Qdrant",
            "Puerta D-MEM (sintesis de creencias)",
            "Backend PostgreSQL para el framework",
            "Integracion del servidor MCP (monolito importa del framework)",
        ]),
        ("Fase 3 (PLANIFICADA)", [
            "pip install soul-framework (publicacion en PyPI)",
            "Documentacion publica y tutoriales",
            "Versionado de API y garantias de estabilidad",
            "Auditoria de seguridad y endurecimiento",
            "Benchmarks vs Mem0/Letta/Hindsight",
            "Contribuciones de la comunidad",
        ]),
    ]

    for phase_name, items in phases:
        add_styled_heading(doc, phase_name, level=2)
        for item in items:
            p = doc.add_paragraph(item, style='List Bullet')
            p.runs[0].font.size = Pt(10)

    # ── Nota al pie ──
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Generado por JARVIS (Equipo SEAL) — Abril 2026")
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x9e, 0x9e, 0x9e)

    # ── Guardar ──
    doc.save(str(DOCX_PATH))
    print(f"Documento guardado: {DOCX_PATH}")
    print(f"Paginas (estimado): ~25")
    print(f"Imagenes generadas: {len(list(IMG_DIR.glob('*.png')))}")

    return str(DOCX_PATH)


if __name__ == "__main__":
    path = build_document()
    print(f"\nListo! Abrir: {path}")
