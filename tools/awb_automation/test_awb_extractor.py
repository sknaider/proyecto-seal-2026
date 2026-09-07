#!/usr/bin/env python3
"""Tests para awb_extractor.py — sin necesidad de PDF real ni API key."""
import json
import sys
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).parent))
import awb_extractor as ext


# ── Test 1: constantes correctas ──────────────────────────────────────────────

def test_auto_fields():
    assert ext.AUTO_FIELDS == ["razon_social", "ruc", "awb", "monto_awb", "fecha_etd"]
    print("PASS test_auto_fields")

def test_manual_fields():
    assert "nro_dam" in ext.MANUAL_FIELDS
    assert "canal" in ext.MANUAL_FIELDS
    assert "estado" in ext.MANUAL_FIELDS
    assert "pago_hdt" in ext.MANUAL_FIELDS
    assert "c_por_cobrar" in ext.MANUAL_FIELDS
    assert "n_operacion" in ext.MANUAL_FIELDS
    print("PASS test_manual_fields")

def test_no_overlap():
    overlap = set(ext.AUTO_FIELDS) & set(ext.MANUAL_FIELDS)
    assert not overlap, f"Campos duplicados: {overlap}"
    print("PASS test_no_overlap")

def test_headers_cover_all_fields():
    all_fields = ext.AUTO_FIELDS + ext.MANUAL_FIELDS
    for f in all_fields:
        assert f in ext.HEADERS, f"Campo sin header: {f}"
    print("PASS test_headers_cover_all_fields")


# ── Test 2: create_excel ──────────────────────────────────────────────────────

def test_create_excel_basic():
    facturas = [
        {
            "razon_social": "IMPORTACIONES SAC",
            "ruc": "20123456789",
            "awb": "083-12345678",
            "monto_awb": "1500.00",
            "fecha_etd": "10/05/2026",
        }
    ]
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out = Path(f.name)
    ext.create_excel(facturas, out)
    assert out.exists()
    assert out.stat().st_size > 0

    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    # Row 1 = headers (Item + fields), Row 2 = data (no sub-header)
    assert ws.cell(row=1, column=1).value == "Item"
    assert ws.cell(row=1, column=2).value == "Razón Social"
    assert ws.cell(row=2, column=1).value == 1        # Item secuencial
    assert ws.cell(row=2, column=2).value == "IMPORTACIONES SAC"
    assert ws.cell(row=2, column=3).value == "20123456789"
    assert ws.cell(row=2, column=4).value == "083-12345678"
    # nro_dam es campo manual (col 7 = Item+5auto+1manual) → vacío
    assert ws.cell(row=2, column=7).value in (None, "")
    out.unlink()
    print("PASS test_create_excel_basic")

def test_create_excel_multiple_facturas():
    facturas = [
        {"razon_social": f"EMPRESA {i}", "ruc": f"201234567{i:02d}",
         "awb": f"083-{i:08d}", "monto_awb": str(i * 100), "fecha_etd": "01/05/2026"}
        for i in range(5)
    ]
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out = Path(f.name)
    ext.create_excel(facturas, out)
    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    # Row 2-6 = 5 facturas (fila 1 = header)
    assert ws.cell(row=6, column=1).value == 5        # Item 5
    assert ws.cell(row=6, column=2).value == "EMPRESA 4"
    out.unlink()
    print("PASS test_create_excel_multiple_facturas")

def test_create_excel_null_fields():
    facturas = [{"razon_social": None, "ruc": None, "awb": "083-99999999",
                 "monto_awb": None, "fecha_etd": None}]
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out = Path(f.name)
    ext.create_excel(facturas, out)
    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    # Col 4 = AWB (Item + razon_social + ruc + awb)
    assert ws.cell(row=2, column=4).value == "083-99999999"
    out.unlink()
    print("PASS test_create_excel_null_fields")


# ── Test 3: image_to_base64 ───────────────────────────────────────────────────

def test_image_to_base64():
    import base64
    content = b"fake image data"
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        f.write(content)
        tmp = Path(f.name)
    result = ext.image_to_base64(tmp)
    assert result == base64.standard_b64encode(content).decode("utf-8")
    tmp.unlink()
    print("PASS test_image_to_base64")


# ── Test 4: JSON parsing en extract_invoices_from_page ────────────────────────

def test_extract_parses_clean_json():
    mock_client = MagicMock()
    mock_response = MagicMock()
    mock_response.content = [MagicMock(text=json.dumps({
        "facturas": [
            {"razon_social": "TEST SA", "ruc": "20000000001",
             "awb": "083-00000001", "monto_awb": "500", "fecha_etd": "01/05/2026"}
        ]
    }))]
    mock_client.messages.create.return_value = mock_response

    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        f.write(b"\x89PNG\r\n\x1a\n" + b"\x00" * 100)
        img = Path(f.name)

    result = ext.extract_invoices_from_page(mock_client, img, page_num=1)
    assert len(result) == 1
    assert result[0]["razon_social"] == "TEST SA"
    assert result[0]["awb"] == "083-00000001"
    img.unlink()
    print("PASS test_extract_parses_clean_json")

def test_extract_handles_markdown_wrapper():
    mock_client = MagicMock()
    raw = '```json\n{"facturas": [{"razon_social": "ACME", "ruc": "20999999999", "awb": "111-22222222", "monto_awb": "200", "fecha_etd": "02/05/2026"}]}\n```'
    mock_response = MagicMock()
    mock_response.content = [MagicMock(text=raw)]
    mock_client.messages.create.return_value = mock_response

    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        f.write(b"\x89PNG\r\n\x1a\n" + b"\x00" * 100)
        img = Path(f.name)

    result = ext.extract_invoices_from_page(mock_client, img, page_num=1)
    assert len(result) == 1
    assert result[0]["razon_social"] == "ACME"
    img.unlink()
    print("PASS test_extract_handles_markdown_wrapper")

def test_extract_handles_bad_json():
    mock_client = MagicMock()
    mock_response = MagicMock()
    mock_response.content = [MagicMock(text="No puedo procesar esta imagen")]
    mock_client.messages.create.return_value = mock_response

    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        f.write(b"\x89PNG\r\n\x1a\n" + b"\x00" * 100)
        img = Path(f.name)

    result = ext.extract_invoices_from_page(mock_client, img, page_num=1)
    assert result == []
    img.unlink()
    print("PASS test_extract_handles_bad_json")


# ── Runner ────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    tests = [
        test_auto_fields,
        test_manual_fields,
        test_no_overlap,
        test_headers_cover_all_fields,
        test_create_excel_basic,
        test_create_excel_multiple_facturas,
        test_create_excel_null_fields,
        test_image_to_base64,
        test_extract_parses_clean_json,
        test_extract_handles_markdown_wrapper,
        test_extract_handles_bad_json,
    ]
    failed = 0
    for t in tests:
        try:
            t()
        except Exception as e:
            print(f"FAIL {t.__name__}: {e}")
            failed += 1
    print(f"\n{'='*40}")
    print(f"Resultados: {len(tests)-failed}/{len(tests)} tests pasados")
    if failed:
        sys.exit(1)
