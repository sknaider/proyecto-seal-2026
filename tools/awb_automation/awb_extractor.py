#!/usr/bin/env python3
"""
AWB Invoice Extractor
Convierte PDFs escaneados de facturas AWB → Excel con campos extraídos via Claude Haiku Vision.

Uso:
    python3 awb_extractor.py documento.pdf
    python3 awb_extractor.py documento.pdf --output resultado.xlsx
    python3 awb_extractor.py documento.pdf --dry-run   # muestra JSON, no genera Excel
"""
from __future__ import annotations

import argparse
import base64
import json
import os
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import anthropic
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image

# ── Config ───────────────────────────────────────────────────────────────────

MODEL = "claude-haiku-4-5-20251001"
MAX_TOKENS = 4096
DPI = 200  # resolución de conversión PDF→imagen

# Campos que Claude extrae automáticamente del PDF
AUTO_FIELDS = [
    "razon_social",
    "ruc",
    "awb",
    "monto_awb",
    "fecha_etd",
]

# Campos que el usuario llena manualmente (quedan vacíos en Excel)
MANUAL_FIELDS = [
    "nro_dam",
    "canal",
    "estado",
    "pago_hdt",
    "c_por_cobrar",
    "n_operacion",
]

HEADERS = {
    "razon_social": "Razón Social",
    "ruc": "RUC",
    "awb": "AWB",
    "monto_awb": "Monto AWB",
    "fecha_etd": "Fecha ETD",
    "nro_dam": "Nº DAM",
    "canal": "Canal",
    "estado": "Estado",
    "pago_hdt": "Pago HDT",
    "c_por_cobrar": "C. por Cobrar",
    "n_operacion": "N. Operación",
}

EXTRACTION_PROMPT = """Analiza esta imagen de un documento de aduanas/logística peruano y extrae TODAS las facturas AWB que encuentres.

Para cada factura AWB en la imagen, extrae SOLO estos 5 campos:
- razon_social: Nombre/razón social del importador o exportador
- ruc: RUC del importador/exportador (11 dígitos)
- awb: Número de AWB (Air Waybill) o guía aérea
- monto_awb: Monto total de cargos del AWB (número, sin símbolo de moneda)
- fecha_etd: Fecha de vuelo o ETD (formato DD/MM/YYYY si es posible)

Si un campo no está visible o no existe en el documento, usa null.
Si hay múltiples facturas AWB en la imagen, devuelve un array con todas.

Responde SOLO con JSON válido, sin texto adicional:
{
  "facturas": [
    {
      "razon_social": "...",
      "ruc": "...",
      "awb": "...",
      "monto_awb": "...",
      "fecha_etd": "..."
    }
  ]
}"""


# ── PDF → imágenes ────────────────────────────────────────────────────────────

def pdf_to_images(pdf_path: Path, dpi: int = DPI) -> list[Path]:
    """Convierte PDF a lista de imágenes PNG via pdftoppm."""
    with tempfile.TemporaryDirectory() as tmpdir:
        prefix = Path(tmpdir) / "page"
        result = subprocess.run(
            ["pdftoppm", "-r", str(dpi), "-png", str(pdf_path), str(prefix)],
            capture_output=True, text=True,
        )
        if result.returncode != 0:
            raise RuntimeError(f"pdftoppm falló: {result.stderr}")

        pages = sorted(Path(tmpdir).glob("page-*.png"))
        if not pages:
            pages = sorted(Path(tmpdir).glob("page*.png"))

        # Copiar a directorio temporal persistente para retorno
        out_dir = Path(tempfile.mkdtemp(prefix="awb_pages_"))
        out_pages = []
        for p in pages:
            dest = out_dir / p.name
            dest.write_bytes(p.read_bytes())
            out_pages.append(dest)
        return out_pages


def image_to_base64(img_path: Path) -> str:
    """Convierte imagen a base64 para API de Claude."""
    return base64.standard_b64encode(img_path.read_bytes()).decode("utf-8")


# ── Extracción Claude Vision ──────────────────────────────────────────────────

def extract_invoices_from_page(
    client: anthropic.Anthropic,
    img_path: Path,
    page_num: int,
) -> list[dict]:
    """Envía una página a Claude Haiku Vision y extrae facturas AWB."""
    img_b64 = image_to_base64(img_path)

    # Detectar media type
    suffix = img_path.suffix.lower()
    media_type = "image/png" if suffix == ".png" else "image/jpeg"

    print(f"  Procesando página {page_num}...", end=" ", flush=True)

    response = client.messages.create(
        model=MODEL,
        max_tokens=MAX_TOKENS,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": media_type,
                            "data": img_b64,
                        },
                    },
                    {
                        "type": "text",
                        "text": EXTRACTION_PROMPT,
                    },
                ],
            }
        ],
    )

    raw = response.content[0].text.strip()

    # Limpiar markdown si Claude lo envuelve en ```json
    if raw.startswith("```"):
        lines = raw.split("\n")
        # Eliminar primera línea (```json) y cortar en el siguiente ``` si existe
        content_lines = []
        for line in lines[1:]:
            if line.strip().startswith("```"):
                break
            content_lines.append(line)
        raw = "\n".join(content_lines)

    try:
        data = json.loads(raw)
        facturas = data.get("facturas", [])
        print(f"{len(facturas)} factura(s) encontrada(s)")
        return facturas
    except json.JSONDecodeError as e:
        print(f"ERROR parsing JSON: {e}")
        print(f"  Respuesta raw: {raw[:200]}")
        return []


# ── Excel output ──────────────────────────────────────────────────────────────

def create_excel(facturas: list[dict], output_path: Path) -> None:
    """Genera el Excel con los campos extraídos y manuales vacíos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facturas AWB"

    all_fields = AUTO_FIELDS + MANUAL_FIELDS
    # Item es la primera columna — número secuencial
    header_labels = ["Item"] + [HEADERS[f] for f in all_fields]

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    item_fill = PatternFill("solid", fgColor="404040")    # gris oscuro = item
    auto_fill = PatternFill("solid", fgColor="1F4E79")    # azul oscuro = auto
    manual_fill = PatternFill("solid", fgColor="833C00")  # naranja = manual
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Fila de headers (fila 1, sin sub-headers)
    for col_idx, label in enumerate(header_labels, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.alignment = center
        cell.border = border
        if col_idx == 1:
            cell.fill = item_fill
        else:
            field = all_fields[col_idx - 2]
            cell.fill = manual_fill if field in MANUAL_FIELDS else auto_fill

    # Datos (empiezan en fila 2)
    for row_idx, factura in enumerate(facturas, start=2):
        # Columna Item
        item_cell = ws.cell(row=row_idx, column=1, value=row_idx - 1)
        item_cell.border = border
        item_cell.alignment = center

        for col_idx, field in enumerate(all_fields, start=2):
            value = factura.get(field, "") if field in AUTO_FIELDS else ""
            cell = ws.cell(row=row_idx, column=col_idx, value=value if value else "")
            cell.border = border
            if field in MANUAL_FIELDS:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")  # amarillo claro

    # Anchos de columna
    col_widths = {
        "razon_social": 35, "ruc": 14, "awb": 18, "monto_awb": 14,
        "fecha_etd": 14, "nro_dam": 18, "canal": 12, "estado": 18,
        "pago_hdt": 14, "c_por_cobrar": 16, "n_operacion": 16,
    }
    ws.column_dimensions["A"].width = 6  # Item
    for col_idx, field in enumerate(all_fields, start=2):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = col_widths.get(field, 16)

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    wb.save(output_path)


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    parser = argparse.ArgumentParser(description="AWB Invoice Extractor — SEAL")
    parser.add_argument("pdf", type=Path, help="PDF de facturas AWB escaneado")
    parser.add_argument("--output", "-o", type=Path, help="Archivo Excel de salida")
    parser.add_argument("--dry-run", action="store_true", help="Solo muestra JSON, no genera Excel")
    parser.add_argument("--dpi", type=int, default=DPI, help=f"DPI conversión (default: {DPI})")
    args = parser.parse_args()

    if not args.pdf.exists():
        print(f"ERROR: No se encuentra el archivo: {args.pdf}")
        sys.exit(1)

    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        print("ERROR: ANTHROPIC_API_KEY no está definida")
        sys.exit(1)

    client = anthropic.Anthropic(api_key=api_key)

    output_path = args.output or args.pdf.with_suffix(".xlsx")

    print(f"\nAWB Extractor — SEAL")
    print(f"PDF: {args.pdf}")
    print(f"Salida: {output_path}\n")

    # 1. Convertir PDF a imágenes
    print("Convirtiendo PDF a imágenes...")
    try:
        pages = pdf_to_images(args.pdf, dpi=args.dpi)
    except RuntimeError as e:
        print(f"ERROR: {e}")
        sys.exit(1)
    print(f"  {len(pages)} página(s) detectada(s)\n")

    # 2. Extraer facturas de cada página
    all_facturas: list[dict] = []
    for i, page_path in enumerate(pages, start=1):
        facturas = extract_invoices_from_page(client, page_path, i)
        for f in facturas:
            f["_pagina"] = i
        all_facturas.extend(facturas)

    # Limpiar imágenes temporales
    for p in pages:
        try:
            p.unlink()
            p.parent.rmdir()
        except Exception:
            pass

    print(f"\nTotal: {len(all_facturas)} factura(s) AWB extraída(s)")

    if not all_facturas:
        print("No se encontraron facturas. Verifica que el PDF sea legible.")
        sys.exit(0)

    if args.dry_run:
        print("\n--- JSON extraído ---")
        print(json.dumps(all_facturas, ensure_ascii=False, indent=2))
        return

    # 3. Generar Excel
    create_excel(all_facturas, output_path)
    print(f"Excel generado: {output_path}")
    print(f"\nCampos AUTO completados: Razón Social, RUC, AWB, Monto AWB, Fecha ETD")
    print(f"Campos MANUAL (vacíos): Nº DAM, Canal, Estado, Pago HDT, C. por Cobrar, N. Operación")


if __name__ == "__main__":
    main()
